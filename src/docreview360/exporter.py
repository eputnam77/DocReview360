"""Export documents to various formats."""

from __future__ import annotations

from html import escape
from io import BytesIO
from typing import Any, Callable, Iterable, Optional
from zipfile import ZipFile

from .documents import DocumentModel, sanitize_content
from .guardrails import (
    GuardrailDiagnostics,
    GuardrailThresholds,
    evaluate_guardrails,
    load_guardrail_thresholds,
)
from .logs import log_event
from .settings_manager import load_settings

# Customisable guardrails checker; return True if content is invalid
guardrails_checker: Optional[Callable[[str], bool]] = None
guardrails_diagnostics: Optional[dict[str, Any]] = None
guardrail_thresholds: GuardrailThresholds | None = None


class ExportSizeLimitError(ValueError):
    """Raised when an export exceeds the configured size ceiling."""

    def __init__(
        self,
        document_id: str,
        fmt: str,
        limit: int,
        size: int,
    ) -> None:
        self.document_id = document_id
        self.format = fmt
        self.limit = limit
        self.size = size
        super().__init__(
            f"{fmt.upper()} export for document {document_id} exceeded the"
            f" {limit} byte limit (generated {size} bytes)"
        )


class _SizeLimitedBuffer(BytesIO):
    """Bytes buffer that tracks whether a configured size limit was exceeded."""

    def __init__(self, limit: int) -> None:
        super().__init__()
        self._limit = limit
        self._exceeded = False

    @property
    def exceeded(self) -> bool:
        return self._exceeded

    def _check(self, prospective_size: int) -> None:
        if prospective_size > self._limit:
            self._exceeded = True

    def write(self, b: bytes | bytearray | memoryview) -> int:  # type: ignore[override]
        if self.closed:
            try:
                return super().write(b)
            except ValueError:
                return 0
        self._check(self.tell() + len(b))
        return super().write(b)

    def seek(self, offset: int, whence: int = 0) -> int:  # type: ignore[override]
        if self.closed:
            try:
                return super().seek(offset, whence)
            except ValueError:
                return 0
        position = super().seek(offset, whence)
        self._check(position)
        return position

    def tell(self) -> int:  # type: ignore[override]
        if self.closed:
            try:
                return super().tell()
            except ValueError:
                return 0
        return super().tell()


_FORMAT_ALIASES: dict[str, str] = {
    "md": "markdown",
    "htm": "html",
}


DEFAULT_EXPORT_SIZE_LIMITS: dict[str, int] = {
    "docx": 5 * 1024 * 1024,
    "pdf": 5 * 1024 * 1024,
}


def _load_export_limits() -> dict[str, int | None]:
    """Return export size ceilings keyed by format with sane defaults."""

    settings = load_settings()
    exports_cfg = settings.get("exports", {}) if isinstance(settings, dict) else {}
    limits_cfg = exports_cfg.get("size_limits", {})
    limits: dict[str, int | None] = {}
    if isinstance(limits_cfg, dict):
        for key, raw in limits_cfg.items():
            fmt = str(key).lower()
            try:
                value = int(raw)
            except (TypeError, ValueError):
                continue
            if value <= 0:
                limits[fmt] = None
            else:
                limits[fmt] = value
    for fmt, default in DEFAULT_EXPORT_SIZE_LIMITS.items():
        limits.setdefault(fmt, default)
    return limits


def _get_export_size_limit(fmt: str) -> int | None:
    """Return the configured export size ceiling for ``fmt`` (or ``None``)."""

    limits = _load_export_limits()
    return limits.get(fmt)


def _canonical_format(fmt: str) -> str:
    """Return the canonical format key for ``fmt``."""

    return _FORMAT_ALIASES.get(fmt, fmt)


def _raise_export_size_error(
    document: DocumentModel,
    fmt: str,
    limit: int,
    size: int,
) -> None:
    """Log and raise an :class:`ExportSizeLimitError`."""

    log_event(
        "export_size_limit_exceeded",
        entity_type="document",
        entity_id=document.id,
        severity="warning",
        fmt=fmt,
        size=size,
        limit=limit,
    )
    raise ExportSizeLimitError(document.id, fmt, limit, size)


def _make_buffer(
    fmt: str,
    limit: int | None = None,
) -> tuple[BytesIO, int | None]:
    """Return a buffer for export with a configured size ceiling."""

    if limit is None:
        limit = _get_export_size_limit(fmt)
    if limit is None:
        return BytesIO(), None

    return _SizeLimitedBuffer(limit), limit


def _ensure_size_within_limit(
    document: DocumentModel,
    fmt: str,
    limit: int | None,
    size: int,
) -> None:
    if limit is None:
        return
    if size > limit:
        _raise_export_size_error(document, fmt, limit, size)


def set_guardrails(
    checker: Callable[[str], bool] | None,
    *,
    thresholds: GuardrailThresholds | None = None,
) -> None:
    """Register a guardrails validation function and thresholds."""

    global guardrails_checker, guardrail_thresholds
    guardrails_checker = checker
    guardrail_thresholds = thresholds or load_guardrail_thresholds()


def get_guardrails_diagnostics() -> Optional[dict[str, Any]]:
    """Return the most recent guardrails diagnostics payload."""

    return guardrails_diagnostics.copy() if guardrails_diagnostics else None


def _record_guardrails_diagnostics(text: str) -> GuardrailDiagnostics:
    global guardrails_diagnostics

    thresholds = guardrail_thresholds or load_guardrail_thresholds()
    diagnostics = evaluate_guardrails(text, thresholds=thresholds)
    guardrails_diagnostics = diagnostics.asdict()
    return diagnostics


def _render_template(template: str, document: DocumentModel) -> str:
    """Return ``template`` formatted with document attributes.

    Supported placeholders are ``id`` and ``number``. Unknown placeholders
    raise :class:`ValueError` to surface configuration issues early.
    """

    try:
        # ``document.number`` is an ``int`` and ``0`` is a valid value.  Using
        # ``or`` would treat ``0`` as falsy and substitute an empty string,
        # yielding filenames like ``doc-.md``.  Explicitly check for ``None`` so
        # that ``0`` is preserved.
        number = document.number if document.number is not None else ""
        base = template.format(id=document.id, number=number)
    except (KeyError, IndexError) as exc:  # pragma: no cover - defensive
        # ``str.format`` raises ``KeyError`` for unknown named fields and
        # ``IndexError`` for stray positional fields such as ``"{0}"``.  The
        # latter previously bubbled up as an ``IndexError`` which callers were
        # not expecting.  Surface both as a ``ValueError`` so users receive a
        # clear, consistent error regardless of placeholder style.
        raise ValueError(f"Unknown placeholder {exc.args[0]!r} in filename template") from exc
    # Avoid path traversal via placeholders
    safe = base.replace("/", "_").replace("\\", "_")
    # Strip newline and null characters which could lead to confusing or
    # truncated filenames on some file systems.
    for ch in ("\n", "\r", "\x00"):
        safe = safe.replace(ch, "_")
    return safe


def export_document(  # noqa: C901
    document: DocumentModel,
    fmt: str,
    *,
    name_template: str | None = None,
) -> tuple[bytes, str, str]:
    """Convert a document into ``fmt`` and return bytes, media type and filename."""
    if not isinstance(fmt, str):
        raise TypeError("fmt must be a string")
    fmt = fmt.lower()
    canonical_fmt = _canonical_format(fmt)
    limit = _get_export_size_limit(canonical_fmt)
    sanitized_content = sanitize_content(document.content)
    diagnostics = _record_guardrails_diagnostics(sanitized_content)
    if diagnostics.blocked:
        raise ValueError("Content failed critic guardrails validation")
    if guardrails_checker is not None:
        if guardrails_checker(sanitized_content):
            raise ValueError("Content failed guardrails validation")
    elif sanitized_content != document.content:
        raise ValueError("Content failed built-in guardrails validation")
    template = name_template or "{id}"
    base = _render_template(template, document)
    if fmt in ("markdown", "md"):
        data = sanitized_content.encode()
        media = "text/markdown"
        name = f"{base}.md"
        _ensure_size_within_limit(document, canonical_fmt, limit, len(data))
    elif fmt in ("html", "htm"):
        body = escape(sanitized_content)
        html_doc = f"<html><body><pre>{body}</pre></body></html>"
        data = html_doc.encode()
        media = "text/html"
        name = f"{base}.html"
        _ensure_size_within_limit(document, canonical_fmt, limit, len(data))
    elif fmt == "json":
        data = document.model_dump_json().encode()
        media = "application/json"
        name = f"{base}.json"
        _ensure_size_within_limit(document, canonical_fmt, limit, len(data))
    elif fmt == "docx":
        import re

        from docx import Document
        from docx.oxml import OxmlElement

        docx_doc = Document()
        for line in document.content.splitlines():
            m = re.match(r"^(#+)\s+(.*)", line)
            if m:
                level = len(m.group(1))
                text = m.group(2)
                p = docx_doc.add_heading("", level=level)
                line_text = text
            else:
                p = docx_doc.add_paragraph()
                line_text = line

            idx = 0
            for match in re.finditer(r"\[(INS|DEL):(.*?)\]", line_text):
                if match.start() > idx:
                    p.add_run(line_text[idx : match.start()])
                tag, txt = match.group(1).lower(), match.group(2)
                change = OxmlElement(f"w:{tag}")
                r = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.text = txt
                r.append(t)
                change.append(r)
                p._element.append(change)
                idx = match.end()
            if idx < len(line_text):
                p.add_run(line_text[idx:])

        buffer, buffer_limit = _make_buffer(canonical_fmt, limit)
        docx_doc.save(buffer)
        if isinstance(buffer, _SizeLimitedBuffer) and buffer_limit is not None and buffer.exceeded:
            _raise_export_size_error(document, canonical_fmt, buffer_limit, buffer.tell())
        data = buffer.getvalue()
        _ensure_size_within_limit(document, canonical_fmt, buffer_limit, len(data))
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        name = f"{base}.docx"
    elif fmt == "pdf":
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
        except Exception:  # pragma: no cover - optional dependency
            buffer, buffer_limit = _make_buffer(canonical_fmt, limit)
            buffer.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF")
            if (
                isinstance(buffer, _SizeLimitedBuffer)
                and buffer_limit is not None
                and buffer.exceeded
            ):
                _raise_export_size_error(document, canonical_fmt, buffer_limit, buffer.tell())
            data = buffer.getvalue()
            _ensure_size_within_limit(document, canonical_fmt, buffer_limit, len(data))
        else:
            buffer, buffer_limit = _make_buffer(canonical_fmt, limit)
            c = canvas.Canvas(buffer, pagesize=letter)
            text = c.beginText(40, 750)
            for line in document.content.splitlines():
                text.textLine(line)
            c.drawText(text)
            c.save()
            if (
                isinstance(buffer, _SizeLimitedBuffer)
                and buffer_limit is not None
                and buffer.exceeded
            ):
                _raise_export_size_error(document, canonical_fmt, buffer_limit, buffer.tell())
            data = buffer.getvalue()
            _ensure_size_within_limit(document, canonical_fmt, buffer_limit, len(data))
        media = "application/pdf"
        name = f"{base}.pdf"
    else:
        raise ValueError(f"Unsupported format: {fmt}")
    return data, media, name


def preview_document(document: DocumentModel, fmt: str, *, length: int = 30) -> str:
    """Return a human-readable preview of a document export."""
    if not isinstance(length, int):
        raise TypeError("length must be an integer")
    if length < 0:
        raise ValueError("length must be non-negative")
    data, media, _ = export_document(document, fmt)
    # For textual formats (``text/*`` or JSON) show the exported representation.
    # The previous implementation returned the original document content for
    # ``application/json`` which was misleading when callers expected to preview
    # the JSON payload.
    if media.startswith("text/") or media == "application/json":
        text = data.decode()
    else:
        text = document.content
    return text[:length]


def export_batch(
    documents: Iterable[DocumentModel],
    fmt: str,
    *,
    name_template: str | None = None,
) -> bytes:
    """Export multiple documents of the same format as a ZIP archive."""
    buffer = BytesIO()
    fmt_lower = fmt.lower()
    canonical_fmt = _canonical_format(fmt_lower)
    limit = _get_export_size_limit(canonical_fmt)
    total_size = 0
    with ZipFile(buffer, "w") as zf:
        for doc in documents:
            data, _, name = export_document(doc, fmt, name_template=name_template)
            _ensure_size_within_limit(doc, canonical_fmt, limit, len(data))
            total_size += len(data)
            _ensure_size_within_limit(doc, canonical_fmt, limit, total_size)
            zf.writestr(name, data)
    return buffer.getvalue()
