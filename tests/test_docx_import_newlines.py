from io import BytesIO

import docx
import pytest

from docreview360 import documents


def test_extract_docx_text_preserves_newlines() -> None:
    """DOCX paragraph boundaries should become real newlines."""
    doc = docx.Document()
    doc.add_paragraph("First")
    doc.add_paragraph("Second")
    buf = BytesIO()
    doc.save(buf)
    text = documents._extract_docx_text(buf.getvalue())
    assert "First\nSecond" in text
    assert "First\\nSecond" not in text


def test_import_document_rejects_invalid_utf8() -> None:
    with pytest.raises(ValueError):
        documents.import_document("note.txt", b"\xff\xfe\xfa")
